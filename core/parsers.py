"""
core/parsers.py
---------------
Parsers de documentos: extrai texto bruto de diferentes formatos.

Formatos suportados:
    .pdf   – via PyMuPDF (fitz)
    .docx  – via python-docx
    .txt   – leitura direta
    .md    – leitura direta
    .xlsx  – via openpyxl (concatena células)

Uso:
    from core.parsers import parse_document
    text = parse_document(Path("relatorio.pdf"))
"""

from __future__ import annotations

from pathlib import Path


def parse_document(path: Path) -> str:
    """Retorna o texto extraído do documento em *path*.

    Raises:
        ValueError: extensão não suportada.
        FileNotFoundError: arquivo inexistente.
    """
    if not path.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {path}")

    ext = path.suffix.lower()

    match ext:
        case ".pdf":
            return _parse_pdf(path)
        case ".docx":
            return _parse_docx(path)
        case ".txt" | ".md":
            return _parse_text(path)
        case ".xlsx":
            return _parse_xlsx(path)
        case _:
            raise ValueError(f"Extensão não suportada: {ext}")


# ---------------------------------------------------------------------------
# Parsers internos
# ---------------------------------------------------------------------------

def _parse_pdf(path: Path) -> str:
    """Extrai texto de PDF página a página usando PyMuPDF."""
    import fitz  # PyMuPDF

    doc = fitz.open(str(path))
    pages: list[str] = []
    for page in doc:
        pages.append(page.get_text())
    doc.close()
    return "\n".join(pages)


def _parse_docx(path: Path) -> str:
    """Extrai texto de DOCX parágrafo a parágrafo."""
    from docx import Document

    doc = Document(str(path))
    return "\n".join(para.text for para in doc.paragraphs if para.text.strip())


def _parse_text(path: Path) -> str:
    """Lê arquivos de texto plano / markdown."""
    return path.read_text(encoding="utf-8", errors="replace")


def _parse_xlsx(path: Path) -> str:
    """Extrai texto de planilhas Excel concatenando células não nulas."""
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True)
    lines: list[str] = []
    for sheet in wb.worksheets:
        lines.append(f"[Planilha: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            row_text = "\t".join(str(c) for c in row if c is not None)
            if row_text.strip():
                lines.append(row_text)
    return "\n".join(lines)
